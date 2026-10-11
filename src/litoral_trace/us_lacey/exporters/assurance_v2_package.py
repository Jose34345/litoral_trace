"""Immutable Assurance V2 packages, deterministic artifact capture and tenant store.

Snapshots are persisted using the existing tenant/RLS-owned append-only operation
event table while V2 schema ownership remains with the integration agent.
No outbound government submission is implied by a generated package.
"""
from __future__ import annotations

import base64
from dataclasses import dataclass
from datetime import datetime, timezone
from hashlib import sha256
import json
from typing import Any, Mapping, Protocol
from uuid import UUID

from litoral_trace.us_lacey.ppq505 import PPQ505_FIELDS_BY_KEY
from litoral_trace.us_lacey.exporters.assurance_v2 import (
    AssuranceCase, GateResult, _canonical, _fingerprint,
    evaluate_filing_gate, export_snapshot_from_case,
)
from litoral_trace.us_lacey.exporters.lacey_excel_builder import build_lacey_excel
from litoral_trace.us_lacey.exporters.lawgs_xml_builder import build_lawgs_xml


class ExportBlocked(ValueError):
    def __init__(self, gate: GateResult):
        self.gate = gate
        super().__init__("Assurance export blocked: " + ", ".join(b.code for b in gate.blockers))


@dataclass(frozen=True)
class AssurancePackage:
    fingerprint: str
    snapshot: dict[str, Any]

    def artifact(self, kind: str) -> bytes:
        self.verify()
        encoded = self.snapshot["artifacts"][kind]
        return base64.b64decode(encoded, validate=True)

    def verify(self) -> None:
        if _fingerprint(self.snapshot) != self.fingerprint:
            raise ValueError("Assurance package fingerprint mismatch")


def issue_assurance_package(case: AssuranceCase, *, generated_at: datetime | None = None) -> AssurancePackage:
    """Raise on uncertainty; capture output BYTES alongside case, provenance and decisions."""
    now = (generated_at or datetime.now(timezone.utc)).astimezone(timezone.utc)
    gate = evaluate_filing_gate(case, now=now)
    if not gate.ready:
        raise ExportBlocked(gate)
    prepared = export_snapshot_from_case(case, now=now)
    xml = build_lawgs_xml(prepared)
    xlsx = build_lacey_excel(prepared).getvalue()
    # PPQ 505 preparation is a structured work product, NOT a filled or
    # submitted federal form. Preserve all 1–18 scope/field entries, unlike
    # the narrower merchandise-only LAWGS serializer.
    ppq505 = _canonical({
        "schema_version": "ppq505-preparation-v2",
        "status": "FILING_READY",
        "submitted_to_agency": False,
        "operation_id": case.operation_id,
        "fields": [
            {"line_reference": f["line_reference"], "key": f["name"],
             "value": f["value"], "evidence_ids": f.get("evidence_ids", []),
             "authority": f["authority"]}
            for f in case.fields if str(f.get("value") or "").strip()
            and f.get("name") in PPQ505_FIELDS_BY_KEY
        ],
    })
    payload = {
        "schema_version": "assurance-v2-package-v1",
        "regulation": "US_LACEY_ACT",
        "status": "FILING_READY",
        "submitted_to_agency": False,
        "generated_at": now.isoformat(),
        "organization_id": case.organization_id,
        "operation_id": case.operation_id,
        "source_set_fingerprint": case.source_set_fingerprint,
        "source_set_revision": case.source_set_revision,
        "case_fingerprint": gate.case_fingerprint,
        "case": case.snapshot(),
        "authorized_export": {
            "operation_id": prepared.operation_id,
            "header": {
                "entry_number": prepared.header.entry_number,
                "importer": prepared.header.importer,
                "estimated_date_of_arrival": prepared.header.estimated_date_of_arrival,
            },
            "plant_lines": [
                {key: getattr(line, key) for key in line.__dataclass_fields__}
                for line in prepared.plant_lines
            ],
        },
        "artifacts": {
            "lawgs_xml": base64.b64encode(xml).decode("ascii"),
            "lacey_excel": base64.b64encode(xlsx).decode("ascii"),
            "ppq505_preparation_json": base64.b64encode(ppq505).decode("ascii"),
        },
    }
    # Includes document hashes, evidence references, reviewer decisions and
    # rule versions because case.snapshot() is a pinned value, not a DB pointer.
    return AssurancePackage(_fingerprint(payload), json.loads(_canonical(payload)))


def draft_assurance_excel(case: AssuranceCase) -> bytes:
    """Watermarked, non-filing draft. No raw candidate becomes an approved value."""
    from openpyxl import Workbook
    from io import BytesIO

    gate = evaluate_filing_gate(case)
    if gate.ready:
        raise ValueError("A ready case should use issue_assurance_package")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "INCOMPLETE DRAFT"
    sheet["A1"] = "DRAFT — NOT READY FOR FILING"
    sheet["A2"] = "This workbook is NOT a PPQ 505 / LAWGS submission."
    sheet["A4"] = "Blocker code"
    sheet["B4"] = "Affected subject"
    for index, blocker in enumerate(gate.blockers, 5):
        sheet.cell(index, 1, blocker.code)
        sheet.cell(index, 2, blocker.subject)
    sheet.column_dimensions["A"].width = 36
    sheet.column_dimensions["B"].width = 52
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


class PackageStore(Protocol):
    def save(self, *, organization_id: int, operation_id: str, package: AssurancePackage) -> None: ...
    def load(self, *, organization_id: int, operation_id: str, fingerprint: str) -> AssurancePackage | None: ...


class MemoryPackageStore:
    """Fixture/local demonstration store. Never configured for production retention."""
    def __init__(self) -> None:
        self._data: dict[tuple[int, str, str], dict[str, Any]] = {}

    def save(self, *, organization_id: int, operation_id: str, package: AssurancePackage) -> None:
        package.verify()
        if package.snapshot["organization_id"] != organization_id or package.snapshot["operation_id"] != operation_id:
            raise PermissionError("Package tenant/operation mismatch")
        key = (organization_id, operation_id, package.fingerprint)
        self._data.setdefault(key, json.loads(_canonical(package.snapshot)))

    def load(self, *, organization_id: int, operation_id: str, fingerprint: str) -> AssurancePackage | None:
        value = self._data.get((organization_id, operation_id, fingerprint))
        if value is None:
            return None
        result = AssurancePackage(fingerprint, json.loads(_canonical(value)))
        result.verify()
        return result


class PostgresOperationEventPackageStore:
    """Scoped RLS-backed adapter; no schema migration or shared controller edits.

    Requires a trusted, tenant-scoped session factory; never accepts an org ID
    from a URL/form/cookie. Existing event CHECK permits PACKAGE_GENERATED.
    Stored JSON is the actual immutable artifact snapshot, not an audit summary.
    """
    def __init__(self, session_factory=None) -> None:
        from litoral_trace.us_lacey.db import get_us_lacey_db_session
        self._session_factory = session_factory or get_us_lacey_db_session

    def _operation(self, session, organization_id: int, operation_id: str):
        from sqlalchemy import select
        from litoral_trace.db.models import UsLaceyOperation
        try:
            oid = UUID(operation_id)
        except (ValueError, TypeError):
            return None
        return session.scalar(select(UsLaceyOperation.id).where(
            UsLaceyOperation.organization_id == organization_id,
            UsLaceyOperation.public_id == oid,
        ))

    def load_for_case_in_transaction(
        self, session, *, organization_id: int, operation_id: str,
        case_fingerprint: str,
    ) -> AssurancePackage | None:
        """Return already-issued immutable bytes for the same authorized case.

        The issuer MUST hold the operation/source-set locks and revalidate
        current authority first. A new source set, revoked evidence or changed
        decision produces a different case fingerprint, never a stale replay.
        """
        from sqlalchemy import select
        from litoral_trace.db.models import UsLaceyOperationEvent

        if len(case_fingerprint) != 64 or any(
            ch not in "0123456789abcdef" for ch in case_fingerprint
        ):
            raise ValueError("Invalid case authority fingerprint")
        internal_id = self._operation(session, organization_id, operation_id)
        if internal_id is None:
            raise PermissionError("Operation not found within tenant")
        records = session.scalars(select(UsLaceyOperationEvent).where(
            UsLaceyOperationEvent.organization_id == organization_id,
            UsLaceyOperationEvent.operation_id == internal_id,
            UsLaceyOperationEvent.event_type == "PACKAGE_GENERATED",
            UsLaceyOperationEvent.details["package"]["case_fingerprint"].as_string()
            == case_fingerprint,
        ).order_by(UsLaceyOperationEvent.id)).all()
        if not records:
            return None
        if len(records) != 1:
            raise ValueError("Duplicate issuance for one authorized case")
        value = records[0].details.get("package")
        if not isinstance(value, dict):
            raise ValueError("Issued package is missing its immutable snapshot")
        package = AssurancePackage(_fingerprint(value), dict(value))
        package.verify()
        if (package.snapshot.get("organization_id") != organization_id
                or package.snapshot.get("operation_id") != operation_id
                or package.snapshot.get("case_fingerprint") != case_fingerprint
                or records[0].event_key != f"assurance_v2:{package.fingerprint}"):
            raise ValueError("Issued package does not match authenticated case")
        return package

    def save_in_transaction(self, session, *, organization_id: int,
                            operation_id: str, package: AssurancePackage) -> None:
        """Stage immutable issuance in the caller's tenant-scoped transaction.

        The trusted issuer must first lock and re-check current document state.
        This method deliberately DOES NOT COMMIT, so a concurrent-source failure
        can roll back both the authorization and the generated artifacts.
        """
        from sqlalchemy import select
        from litoral_trace.db.models import UsLaceyOperationEvent

        package.verify()
        if package.snapshot["organization_id"] != organization_id or package.snapshot["operation_id"] != operation_id:
            raise PermissionError("Package tenant/operation mismatch")
        internal_id = self._operation(session, organization_id, operation_id)
        if internal_id is None:
            raise PermissionError("Operation not found within tenant")
        key = f"assurance_v2:{package.fingerprint}"
        existing = session.scalar(select(UsLaceyOperationEvent).where(
            UsLaceyOperationEvent.organization_id == organization_id,
            UsLaceyOperationEvent.operation_id == internal_id,
            UsLaceyOperationEvent.event_key == key,
        ))
        if existing is not None:
            if existing.details.get("package") != package.snapshot:
                raise ValueError("Immutable package collision")
            return
        session.add(UsLaceyOperationEvent(
            organization_id=organization_id,
            operation_id=internal_id,
            actor_type="USER",
            actor_identity=str(package.snapshot["case"]["export_authorization"]["actor_id"])[:255],
            event_type="PACKAGE_GENERATED",
            event_key=key,
            details={"package_type": "ASSURANCE_V2", "package": package.snapshot},
        ))
        session.flush()

    def save(self, *, organization_id: int, operation_id: str, package: AssurancePackage) -> None:
        from litoral_trace.db.tenant import set_tenant_db_context

        session = self._session_factory()
        try:
            set_tenant_db_context(session, organization_id)
            self.save_in_transaction(
                session, organization_id=organization_id,
                operation_id=operation_id, package=package,
            )
            session.commit()
        except Exception:
            session.rollback()
            raise
        finally:
            session.close()

    def load(self, *, organization_id: int, operation_id: str, fingerprint: str) -> AssurancePackage | None:
        from sqlalchemy import select
        from litoral_trace.db.models import UsLaceyOperationEvent
        from litoral_trace.db.tenant import set_tenant_db_context

        if len(fingerprint) != 64 or any(c not in "0123456789abcdef" for c in fingerprint):
            return None
        session = self._session_factory()
        try:
            set_tenant_db_context(session, organization_id)
            internal_id = self._operation(session, organization_id, operation_id)
            if internal_id is None:
                return None
            row = session.scalar(select(UsLaceyOperationEvent).where(
                UsLaceyOperationEvent.organization_id == organization_id,
                UsLaceyOperationEvent.operation_id == internal_id,
                UsLaceyOperationEvent.event_key == f"assurance_v2:{fingerprint}",
                UsLaceyOperationEvent.event_type == "PACKAGE_GENERATED",
            ))
            if row is None or not isinstance(row.details.get("package"), dict):
                return None
            package = AssurancePackage(fingerprint, dict(row.details["package"]))
            package.verify()
            if package.snapshot.get("organization_id") != organization_id or package.snapshot.get("operation_id") != operation_id:
                return None
            return package
        finally:
            session.close()
