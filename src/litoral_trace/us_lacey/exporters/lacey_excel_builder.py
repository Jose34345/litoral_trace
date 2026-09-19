"""Human-readable XLSX serializer for the shared Phase E export snapshot."""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from litoral_trace.us_lacey.exporters.export_snapshot import LaceyExportSnapshot


_HEADERS = (
    "Line Reference",
    "HTS Number",
    "Entered Value",
    "Component Description",
    "Genus",
    "Species",
    "Country of Harvest",
    "Quantity",
    "Unit",
)


def build_lacey_excel(snapshot: LaceyExportSnapshot) -> BytesIO:
    """Return an in-memory XLSX built from the exact same projection as the XML."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Lacey Act Declaration"

    sheet.append(["Entry Number", snapshot.header.entry_number])
    sheet.append(["Importer", snapshot.header.importer])
    sheet.append(["Estimated Date of Arrival", snapshot.header.estimated_date_of_arrival])
    sheet.append(["Operation ID", snapshot.operation_id])
    sheet.append(["Client Reference", snapshot.client_reference])
    sheet.append([])
    sheet.append(list(_HEADERS))

    for line in snapshot.plant_lines:
        sheet.append(
            [
                line.line_reference,
                line.hts_number,
                line.entered_value,
                line.article_component,
                line.genus,
                line.species,
                line.country_of_harvest,
                line.quantity,
                line.unit,
            ]
        )

    sheet.freeze_panes = "A8"
    if sheet.max_row >= 7:
        sheet.auto_filter.ref = f"A7:I{sheet.max_row}"

    widths = (20, 16, 16, 40, 20, 20, 22, 14, 12)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
