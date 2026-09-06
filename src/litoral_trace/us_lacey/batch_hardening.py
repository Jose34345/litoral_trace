"""Memory/complexity guardrails for one-shipment U.S. Lacey uploads."""
from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
import os
from pathlib import PurePath
from typing import Iterable

from openpyxl import load_workbook

from litoral_trace.assurance.tabular_safety import TabularSafetyError, profile_csv_chunks


DATASET_TOO_LARGE = "DATASET_TOO_LARGE_FOR_SHIPMENT_PIPELINE"
DATASET_TOO_COMPLEX = "DATASET_TOO_COMPLEX_FOR_SHIPMENT_PIPELINE"


@dataclass(frozen=True, slots=True)
class ShipmentSpreadsheetLimits:
    max_bytes: int
    max_rows: int
    max_columns: int
    max_cells: int


@dataclass(frozen=True, slots=True)
class ShipmentSpreadsheetProfile:
    file_kind: str
    size_bytes: int
    rows: int
    columns: int
    cells: int


class ShipmentBatchRejected(ValueError):
    def __init__(self, code: str, safe_message: str):
        super().__init__(safe_message)
        self.code = code
        self.safe_message = safe_message


def _env_int(name: str, default: int, minimum: int, maximum: int) -> int:
    raw = str(os.environ.get(name, default)).strip()
    try:
        value = int(raw)
    except ValueError:
        value = default
    return max(minimum, min(maximum, value))


def shipment_spreadsheet_limits() -> ShipmentSpreadsheetLimits:
    return ShipmentSpreadsheetLimits(
        max_bytes=_env_int("US_LACEY_SHIPMENT_SPREADSHEET_MAX_BYTES", 5 * 1024 * 1024, 128 * 1024, 25 * 1024 * 1024),
        max_rows=_env_int("US_LACEY_SHIPMENT_SPREADSHEET_MAX_ROWS", 5000, 100, 100000),
        max_columns=_env_int("US_LACEY_SHIPMENT_SPREADSHEET_MAX_COLUMNS", 64, 4, 512),
        max_cells=_env_int("US_LACEY_SHIPMENT_SPREADSHEET_MAX_CELLS", 100000, 1000, 2_000_000),
    )


def _bulk_message() -> str:
    return (
        "This spreadsheet appears to contain a bulk or multi-shipment dataset and is too large for one Lacey operation. "
        "Use the Litoral bulk benchmark importer instead; the shipment workspace is reserved for documents belonging to one shipment."
    )


def _reject_for_size(size_bytes: int, limits: ShipmentSpreadsheetLimits) -> None:
    if size_bytes > limits.max_bytes:
        raise ShipmentBatchRejected(DATASET_TOO_LARGE, _bulk_message())


def inspect_csv_chunks_for_shipment(
    chunks: Iterable[bytes],
    *,
    size_bytes: int,
    limits: ShipmentSpreadsheetLimits | None = None,
) -> ShipmentSpreadsheetProfile:
    budget = limits or shipment_spreadsheet_limits()
    _reject_for_size(size_bytes, budget)
    try:
        profile = profile_csv_chunks(
            chunks,
            max_rows=budget.max_rows,
            max_columns=budget.max_columns,
            max_cells=budget.max_cells,
        )
    except TabularSafetyError as exc:
        detail = str(exc)
        if detail.startswith(("CSV_ROW_LIMIT_EXCEEDED", "CSV_COLUMN_LIMIT_EXCEEDED", "CSV_CELL_LIMIT_EXCEEDED")):
            raise ShipmentBatchRejected(DATASET_TOO_COMPLEX, _bulk_message()) from exc
        raise ShipmentBatchRejected("INVALID_SHIPMENT_SPREADSHEET", "The CSV could not be safely interpreted as a shipment spreadsheet.") from exc
    return ShipmentSpreadsheetProfile("CSV", size_bytes, profile.row_count, profile.max_columns, profile.total_cells)


def _inspect_xlsx(content: bytes, budget: ShipmentSpreadsheetLimits) -> ShipmentSpreadsheetProfile:
    rows_total = 0
    max_columns = 0
    cells_total = 0
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ShipmentBatchRejected("INVALID_SHIPMENT_SPREADSHEET", "The XLSX could not be safely opened.") from exc
    try:
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows(values_only=True):
                width = len(row)
                max_columns = max(max_columns, width)
                if width > budget.max_columns:
                    raise ShipmentBatchRejected(DATASET_TOO_COMPLEX, _bulk_message())
                if not any(value not in (None, "") for value in row):
                    continue
                rows_total += 1
                cells_total += width
                if rows_total > budget.max_rows or cells_total > budget.max_cells:
                    raise ShipmentBatchRejected(DATASET_TOO_COMPLEX, _bulk_message())
    finally:
        workbook.close()
    return ShipmentSpreadsheetProfile("XLSX", len(content), rows_total, max_columns, cells_total)


def _inspect_xls(content: bytes, budget: ShipmentSpreadsheetLimits) -> ShipmentSpreadsheetProfile:
    try:
        import xlrd
        workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
    except Exception as exc:
        raise ShipmentBatchRejected("INVALID_SHIPMENT_SPREADSHEET", "The XLS file could not be safely opened.") from exc
    rows_total = 0
    max_columns = 0
    cells_total = 0
    try:
        for sheet_name in workbook.sheet_names():
            sheet = workbook.sheet_by_name(sheet_name)
            rows_total += int(sheet.nrows)
            max_columns = max(max_columns, int(sheet.ncols))
            cells_total += int(sheet.nrows) * int(sheet.ncols)
            if rows_total > budget.max_rows or max_columns > budget.max_columns or cells_total > budget.max_cells:
                raise ShipmentBatchRejected(DATASET_TOO_COMPLEX, _bulk_message())
    finally:
        try:
            workbook.release_resources()
        except Exception:
            pass
    return ShipmentSpreadsheetProfile("XLS", len(content), rows_total, max_columns, cells_total)


def enforce_shipment_document_budget(
    *,
    filename: str,
    content: bytes,
    limits: ShipmentSpreadsheetLimits | None = None,
) -> ShipmentSpreadsheetProfile | None:
    """Fail before Vault/queue work if a spreadsheet is a bulk dataset."""
    suffix = PurePath(str(filename or "")).suffix.lower()
    if suffix not in {".csv", ".xlsx", ".xls"}:
        return None
    budget = limits or shipment_spreadsheet_limits()
    _reject_for_size(len(content), budget)
    if suffix == ".csv":
        return inspect_csv_chunks_for_shipment((content,), size_bytes=len(content), limits=budget)
    if suffix == ".xlsx":
        return _inspect_xlsx(content, budget)
    return _inspect_xls(content, budget)


def enforce_streamed_csv_budget(
    *,
    chunks: Iterable[bytes],
    size_bytes: int,
    limits: ShipmentSpreadsheetLimits | None = None,
) -> ShipmentSpreadsheetProfile:
    """Worker-side guard for files already stored before this hardening shipped."""
    return inspect_csv_chunks_for_shipment(chunks, size_bytes=size_bytes, limits=limits)
