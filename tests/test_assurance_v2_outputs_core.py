"""Assurance V2 regulatory authority, serialization and immutable package tests."""
from __future__ import annotations

from dataclasses import replace
from datetime import datetime, timezone
from io import BytesIO
from types import SimpleNamespace
from uuid import UUID
import xml.etree.ElementTree as ET

import pytest
from openpyxl import load_workbook

from litoral_trace.us_lacey.ppq505 import PPQ505_SHIPMENT_REFERENCE
from litoral_trace.us_lacey.exporters.assurance_v2 import AssuranceCase, evaluate_filing_gate
from litoral_trace.us_lacey.exporters.assurance_v2_package import (
    ExportBlocked, MemoryPackageStore, draft_assurance_excel, issue_assurance_package,
)
from litoral_trace.us_lacey.exporters.export_snapshot import consolidate_export_snapshot

OPERATION_ID = "11111111-2222-3333-4444-555555555555"
NOW = datetime(2026, 10, 9, 15, 0, tzinfo=timezone.utc)
VALUES = {
    PPQ505_SHIPMENT_REFERENCE: {
        "estimated_arrival_date": "2026-10-25",
        "filing_entry_reference": "123-4567890-1",
        "manufacturer_id": "USABC123456",
        "importer_name": "Acme Imports",
        "consignee_name": "Acme Distribution",
        "importer_address": "123 Harbor Road, NY",
        "consignee_address": "456 Main Street, NY",
        "merchandise_description": "Oak furniture",
    },
    "LINE-1": {
        "hts_code": "4407990000",
        "entered_value": "18600",
        "article_component": "Sawn oak boards",
        "genus": "Quercus",
        "species": "alba",
        "country_of_harvest": "United States",
        "plant_quantity": "24.5",
        "metric_unit": "kg",
    },
}


def make_case(*, organization_id: int = 11, approved: bool = True) -> AssuranceCase:
    fields = tuple({
        "line_reference": line,
        "name": name,
        "value": value,
        "authority": "SUPPORTED",
        "document_ids": ["doc-invoice"],
        "evidence_ids": [f"claim:{line}:{name}"],
    } for line, values in VALUES.items() for name, value in values.items())
    case = AssuranceCase(
        organization_id=organization_id,
        operation_id=OPERATION_ID,
        source_set_fingerprint="sset-sha256-001",
        current_source_set_fingerprint="sset-sha256-001",
        source_set_revision=3,
        current_source_set_revision=3,
        ruleset_version="lacey-2026-10",
        documents=({"id": "doc-invoice", "name": "Commercial Invoice",
                    "sha256": "a" * 64, "version": 2, "role": "COMMERCIAL_INVOICE"},),
        lines=("LINE-1",),
        fields=fields,
        identities=({"line_reference": "LINE-1", "supplier_id": "sup-us-1",
                     "product_id": "prod-oak", "sku": "OAK-PANELS-24",
                     "status": "VERIFIED"},),
        rules=({"id": "HTS_APPLICABILITY", "version": "lacey-2026-10",
                "inputs_fingerprint": "ri-fp-01", "status": "PASS", "blocking": True},),
    )
    if approved:
        return authorize(case)
    return case


def authorize(case: AssuranceCase) -> AssuranceCase:
    return replace(case, export_authorization={
        "id": "export-authorization-01", "actor_id": "user:42",
        "authorized_at": NOW.isoformat(), "status": "AUTHORIZED",
        "source_set_fingerprint": case.source_set_fingerprint,
        "case_fingerprint": case.authority_fingerprint(),
    })


def with_field(case: AssuranceCase, name: str, **changes) -> AssuranceCase:
    fields = tuple({**field, **changes} if field["name"] == name and
                   field["line_reference"] == "LINE-1" else field for field in case.fields)
    return authorize(replace(case, fields=fields))


def codes(case: AssuranceCase) -> set[str]:
    return {block.code for block in evaluate_filing_gate(case, now=NOW).blockers}


def test_complete_reviewed_case_exports_same_authorized_values_to_xml_and_excel():
    case = make_case()
    gate = evaluate_filing_gate(case, now=NOW)
    assert gate.ready, gate.blockers
    package = issue_assurance_package(case, generated_at=NOW)
    assert package.snapshot["status"] == "FILING_READY"
    assert package.snapshot["submitted_to_agency"] is False
    from litoral_trace.us_lacey.exporters.lawgs_xml_builder import LAWGS_XML_NAMESPACE
    root = ET.fromstring(package.artifact("lawgs_xml"))
    ns = {"x": LAWGS_XML_NAMESPACE}
    assert root.findtext("./x:merchandise/x:genus", namespaces=ns) == "Quercus"
    assert root.findtext("./x:merchandise/x:enteredValue", namespaces=ns) == "18600"
    workbook = load_workbook(BytesIO(package.artifact("lacey_excel")), read_only=True)
    assert workbook.active["E8"].value == "Quercus"
    assert workbook.active["D8"].value == "Sawn oak boards"
    import json
    prepared_ppq = json.loads(package.artifact("ppq505_preparation_json"))
    assert prepared_ppq["submitted_to_agency"] is False
    assert len(prepared_ppq["fields"]) == sum(map(len, VALUES.values()))
    assert any(f["key"] == "genus" and f["value"] == "Quercus" for f in prepared_ppq["fields"])
    assert package.snapshot["case"]["documents"][0]["sha256"] == "a" * 64
    assert package.snapshot["case"]["export_authorization"]["actor_id"] == "user:42"


def test_unsupported_claim_never_leaks_to_filing_ready():
    bad = with_field(make_case(), "species", authority="CANDIDATE")
    assert "UNSUPPORTED_CLAIM" in codes(bad)
    with pytest.raises(ExportBlocked):
        issue_assurance_package(bad, generated_at=NOW)
    draft = load_workbook(BytesIO(draft_assurance_excel(bad)), read_only=True)
    assert draft.active["A1"].value == "DRAFT — NOT READY FOR FILING"
    assert "CANDIDATE" not in str(draft.active.values)


def test_blocking_conflict_stops_final_export_even_if_other_fields_valid():
    case = make_case()
    bad = authorize(replace(case, exceptions=({
        "id": "ex-1", "field_name": "species", "line_reference": "LINE-1",
        "blocking": True, "status": "OPEN", "candidates": [
            {"value": "alba", "document_id": "doc-invoice", "page": 1},
            {"value": "rubra", "document_id": "doc-invoice", "page": 2},
        ],
    },)))
    assert {"BLOCKING_EXCEPTION", "UNDECIDED_CONFLICT"} <= codes(bad)


@pytest.mark.parametrize("state,valid_until,revoked", [
    ("REVOKED", "2030-01-01T00:00:00Z", "2026-10-01T00:00:00Z"),
    ("ACTIVE", "2026-01-01T00:00:00Z", None),
    ("ACTIVE", "2030-01-01T00:00:00Z", "2026-10-08T00:00:00Z"),
])
def test_revoked_or_expired_memory_cannot_be_current(state, valid_until, revoked):
    case = make_case()
    reuse = {
        "id": "memory-1", "status": state, "valid_until": valid_until,
        "revoked_at": revoked, "verified_by": "user:7",
        "evidence_id": "history:genus", "source_document_hash": "b" * 64,
        "line_reference": "LINE-1", "supplier_id": "sup-us-1", "product_id": "prod-oak",
    }
    updated = replace(case, reuses=(reuse,))
    updated = with_field(updated, "genus", authority="VERIFIED_REUSE",
                         reuse_id="memory-1", evidence_ids=["history:genus"], document_ids=[])
    assert "INVALID_REUSED_EVIDENCE" in codes(updated)


def test_verified_memory_uses_exact_supplier_product_link_and_document_hash():
    case = make_case()
    reuse = {"id": "r1", "status": "VERIFIED", "valid_until": "2030-10-01T00:00:00Z",
             "verified_by": "user:7", "evidence_id": "history:genus",
             "source_document_hash": "b" * 64, "line_reference": "LINE-1",
             "supplier_id": "sup-us-1", "product_id": "prod-oak"}
    good = with_field(replace(case, reuses=(reuse,)), "genus", authority="VERIFIED_REUSE",
                      reuse_id="r1", evidence_ids=["history:genus"], document_ids=[])
    assert evaluate_filing_gate(good, now=NOW).ready
    mismatch = authorize(replace(good, reuses=({**reuse, "product_id": "another-product"},)))
    assert "INVALID_REUSED_EVIDENCE" in codes(mismatch)


def test_document_supersession_invalidates_current_gate_but_preserves_historical_package():
    original = make_case()
    package = issue_assurance_package(original, generated_at=NOW)
    stored = MemoryPackageStore()
    stored.save(organization_id=11, operation_id=OPERATION_ID, package=package)
    stale = replace(original, current_source_set_fingerprint="updated-sset", current_source_set_revision=4)
    assert "STALE_SOURCE_SET" in codes(stale)
    with pytest.raises(ExportBlocked):
        issue_assurance_package(stale, generated_at=NOW)
    historical = stored.load(organization_id=11, operation_id=OPERATION_ID, fingerprint=package.fingerprint)
    assert historical.artifact("lawgs_xml") == package.artifact("lawgs_xml")
    assert historical.artifact("lacey_excel") == package.artifact("lacey_excel")


def test_human_decision_must_match_field_actor_and_value():
    case = make_case()
    confirmed = with_field(case, "genus", authority="HUMAN_CONFIRMED", decision_id="d1")
    assert "MISSING_REVIEW_DECISION" in codes(confirmed)
    fixed = authorize(replace(confirmed, decisions=({
        "id": "d1", "field_name": "genus", "line_reference": "LINE-1",
        "value": "Quercus", "action": "CORRECT", "actor_id": "user:42",
        "decided_at": NOW.isoformat(), "reason": "Verified with supplier certificate",
    },)))
    assert evaluate_filing_gate(fixed, now=NOW).ready
    pkg = issue_assurance_package(fixed, generated_at=NOW)
    assert pkg.snapshot["case"]["decisions"][0]["reason"] == "Verified with supplier certificate"


def test_package_fingerprint_and_tenant_isolation():
    case = make_case()
    package = issue_assurance_package(case, generated_at=NOW)
    store = MemoryPackageStore()
    store.save(organization_id=11, operation_id=OPERATION_ID, package=package)
    assert store.load(organization_id=12, operation_id=OPERATION_ID, fingerprint=package.fingerprint) is None
    assert store.load(organization_id=11, operation_id=OPERATION_ID, fingerprint=package.fingerprint)
    with pytest.raises(PermissionError):
        store.save(organization_id=12, operation_id=OPERATION_ID, package=package)
    package.snapshot["case"]["documents"][0]["sha256"] = "tampered"
    with pytest.raises(ValueError):
        package.verify()


def test_unapproved_proposed_value_no_longer_enters_legacy_snapshot():
    unsupported = SimpleNamespace(
        line_reference="LINE-1", field_name="genus", scope="PLANT_LINE",
        status="REVIEW_REQUIRED", effective_value=None, proposed_value="InventedSpecies",
        source_assurance_document_id=None, source_page=None, source_locator=None,
    )
    detail = SimpleNamespace(public_id=UUID(OPERATION_ID), client_reference="fixture",
                             importer_name="", fields=(unsupported,), plant_declarations=())
    result = consolidate_export_snapshot(detail=detail, evidence_snapshot={})
    assert result.plant_lines[0].genus == ""


def test_current_regulatory_snapshot_adapter_respects_rule_scoped_blocking():
    from litoral_trace.us_lacey.regulatory.assurance_v2_adapter import rules_from_current_assessment
    native = SimpleNamespace(
        status="CURRENT", source_set_fingerprint="sset-sha256-001",
        ruleset_version="lacey-2026-10", input_fingerprint="inputs-sha256",
        payload={"assessments": [
            {"rule_id": "DE_MINIMIS", "status": "INDETERMINATE",
             "review_required": False, "subject_ref": "LINE-1",
             "reason_codes": ["EXEMPTION_NOT_CLAIMED"], "evidence_refs": []},
            {"rule_id": "HTS_APPLICABILITY", "status": "INDETERMINATE",
             "review_required": True, "subject_ref": "LINE-1",
             "reason_codes": ["HTS10_MISSING"], "evidence_refs": []},
        ]},
    )
    rules = rules_from_current_assessment(
        native, current_source_set_fingerprint="sset-sha256-001",
        current_ruleset_version="lacey-2026-10",
    )
    assert rules[0]["blocking"] is False
    assert rules[1]["blocking"] is True
    assert rules[0]["reason_codes"] == ["EXEMPTION_NOT_CLAIMED"]
    assert not rules_from_current_assessment(
        native, current_source_set_fingerprint="new-source-set",
        current_ruleset_version="lacey-2026-10",
    )
