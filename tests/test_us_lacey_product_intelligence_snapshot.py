from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace

from openpyxl import Workbook
import pytest

from litoral_trace.assurance.parsers import ParsedDocument, ParsedTable, SourceLocation
from litoral_trace.product_intelligence.domain import BomIssueSeverity
from litoral_trace.us_lacey.product_intelligence_snapshot import (
    ProductIntelligenceDocumentInput,
    analyze_product_intelligence_documents,
    snapshot_matches_claim,
)
from litoral_trace.us_lacey.source_sets import SourceSetClaim


def _source(filename: str, content: bytes, *, document_id: int = 11) -> ProductIntelligenceDocumentInput:
    return ProductIntelligenceDocumentInput(
        operation_document_id=document_id + 100,
        assurance_document_id=document_id,
        document_id=f"assurance-{document_id}",
        filename=filename,
        source_sha256=("a" * 63) + str(document_id % 10),
        content=content,
    )


def _xlsx_bytes(rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "BOM"
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def test_csv_explicit_bom_builds_ready_snapshot_payload_with_physical_provenance():
    payload = (
        b"SKU,Product,Component,Material,Qty,Weight,UOM\n"
        b"CHAIR-001,Chair,Front leg,Rubberwood,2,0.5,kg\n"
        b"\n"
        b"CHAIR-001,Chair,Seat,Plywood,1,1200,g\n"
    )

    result = analyze_product_intelligence_documents((_source("bom.csv", payload),))

    assert result.status == "READY"
    assert result.unique_sku_count == 1
    assert result.component_count == 2
    assert result.material_count == 2
    assert result.issue_count == 0
    assert result.recognized_bom_table_count == 1
    table = result.payload["sources"][0]["tables"][0]
    assert table["compositions"][0]["sku"] == "CHAIR-001"
    assert table["compositions"][0]["components"][1]["source"]["row"] == 4
    assert table["compositions"][0]["components"][1]["material"]["mass"]["kilograms"] == "1.200"
    assert table["source"]["locator"] == "csv:header_row:1"


def test_xlsx_explicit_bom_builds_ready_snapshot_payload():
    content = _xlsx_bytes(
        [
            ["SKU", "Product", "Component", "Material", "Qty", "Weight", "UOM"],
            ["TABLE-001", "Table", "Top", "Oak", 1, 2, "kg"],
        ]
    )

    result = analyze_product_intelligence_documents((_source("bom.xlsx", content),))

    assert result.status == "READY"
    assert result.unique_sku_count == 1
    assert result.component_count == 1
    assert result.payload["sources"][0]["tables"][0]["source"]["sheet"] == "BOM"


def test_tabular_document_without_required_bom_columns_is_not_applicable():
    payload = b"invoice,amount,currency\nINV-1,100,USD\n"

    result = analyze_product_intelligence_documents((_source("invoice.csv", payload),))

    assert result.status == "NOT_APPLICABLE"
    assert result.recognized_bom_table_count == 0
    assert result.component_count == 0
    assert result.issue_count == 0


def test_mixed_valid_and_invalid_bom_rows_is_partial_and_keeps_usable_composition():
    payload = (
        b"SKU,Product,Component,Material,Qty,Weight,UOM\n"
        b"CHAIR-001,Chair,Leg,Oak,4,1,kg\n"
        b"CHAIR-001,Chair,Seat,Plywood,1,not-a-number,kg\n"
    )

    result = analyze_product_intelligence_documents((_source("bom.csv", payload),))

    assert result.status == "PARTIAL"
    assert result.component_count == 1
    assert result.issue_count == 1
    issue = result.payload["sources"][0]["tables"][0]["issues"][0]
    assert issue["code"] == "INVALID_MASS"
    assert issue["severity"] == BomIssueSeverity.ERROR.value
    assert issue["source"]["row"] == 3


def test_recognized_bom_with_no_usable_rows_is_failed():
    payload = (
        b"SKU,Product,Component,Material,Qty,Weight,UOM\n"
        b",Chair,Leg,Oak,4,1,kg\n"
    )

    result = analyze_product_intelligence_documents((_source("bom.csv", payload),))

    assert result.status == "FAILED"
    assert result.recognized_bom_table_count == 1
    assert result.component_count == 0
    assert result.issue_count == 1


def test_one_parser_failure_plus_one_valid_bom_is_partial_without_discarding_good_source():
    good = (
        b"SKU,Product,Component,Material,Qty,Weight,UOM\n"
        b"CHAIR-001,Chair,Leg,Oak,4,1,kg\n"
    )
    bad_xlsx = b"PK\x03\x04not-a-real-workbook"

    result = analyze_product_intelligence_documents(
        (_source("bom.csv", good, document_id=11), _source("broken.xlsx", bad_xlsx, document_id=12))
    )

    assert result.status == "PARTIAL"
    assert result.component_count == 1
    assert result.issue_count >= 1
    assert any(issue["code"] == "DOCUMENT_PARSE_FAILED" for issue in result.payload["issues"])


def test_snapshot_claim_match_requires_exact_current_finalizing_generation_and_token():
    claim = SourceSetClaim(
        revision_id=9,
        generation=3,
        fingerprint="f" * 64,
        claimed=True,
        reason="CLAIMED",
        claimed_at=SimpleNamespace(),
    )
    revision = SimpleNamespace(
        id=9,
        generation=3,
        source_set_fingerprint="f" * 64,
        status="FINALIZING",
        is_current=True,
        claimed_at=claim.claimed_at,
    )
    assert snapshot_matches_claim(revision, claim) is True

    revision.status = "FINALIZED"
    assert snapshot_matches_claim(revision, claim) is False


def test_snapshot_claim_match_rejects_unclaimed_or_mismatched_claim():
    token = SimpleNamespace()
    revision = SimpleNamespace(
        id=9,
        generation=3,
        source_set_fingerprint="f" * 64,
        status="FINALIZING",
        is_current=True,
        claimed_at=token,
    )
    unclaimed = SourceSetClaim(9, 3, "f" * 64, False, "ALREADY_CLAIMED", token)
    assert snapshot_matches_claim(revision, unclaimed) is False

    wrong = SourceSetClaim(9, 3, "e" * 64, True, "CLAIMED", token)
    assert snapshot_matches_claim(revision, wrong) is False



def test_pdf_explicit_bom_table_builds_ready_product_intelligence(monkeypatch) -> None:
    import litoral_trace.us_lacey.product_intelligence_snapshot as module

    parsed = ParsedDocument(
        file_kind="PDF",
        text="Bill of materials",
        tables=(
            ParsedTable(
                name="page_2_table_1",
                headers=("SKU", "Product", "Component", "Material", "Qty", "Weight", "UOM"),
                rows=(
                    {
                        "SKU": "CHAIR-PDF-1",
                        "Product": "Chair",
                        "Component": "Leg",
                        "Material": "Rubberwood",
                        "Qty": "4",
                        "Weight": "0.5",
                        "UOM": "kg",
                    },
                ),
                source=SourceLocation(
                    page=2,
                    row=1,
                    locator="pdf:page:2;table:1;header_row:1",
                ),
                row_numbers=(2,),
            ),
        ),
        metadata={"page_count": 2},
        ocr_required=False,
    )
    calls: list[str] = []

    def fake_parse(filename: str, content: bytes) -> ParsedDocument:
        calls.append(filename)
        assert content == b"%PDF-explicit-bom"
        return parsed

    monkeypatch.setattr(module, "parse_document", fake_parse)

    result = analyze_product_intelligence_documents(
        (_source("bom.pdf", b"%PDF-explicit-bom"),)
    )

    assert calls == ["bom.pdf"]
    assert result.status == "READY"
    assert result.eligible_document_count == 1
    assert result.recognized_bom_table_count == 1
    assert result.unique_sku_count == 1
    assert result.component_count == 1
    table = result.payload["sources"][0]["tables"][0]
    assert table["source"]["page"] == 2
    assert table["source"]["locator"] == "pdf:page:2;table:1;header_row:1"
    component = table["compositions"][0]["components"][0]
    assert component["source"]["row"] == 2
    assert component["material"]["source"]["row"] == 2


def test_pdf_free_prose_without_explicit_table_remains_not_applicable(monkeypatch) -> None:
    import litoral_trace.us_lacey.product_intelligence_snapshot as module

    parsed = ParsedDocument(
        file_kind="PDF",
        text=(
            "SKU CHAIR-1 uses four rubberwood legs weighing 0.5 kg each. "
            "This prose must never be reconstructed into a BOM."
        ),
        tables=(),
        metadata={"page_count": 1},
        ocr_required=False,
    )

    monkeypatch.setattr(module, "parse_document", lambda _filename, _content: parsed)

    result = analyze_product_intelligence_documents(
        (_source("narrative.pdf", b"%PDF-free-prose"),)
    )

    assert result.status == "NOT_APPLICABLE"
    assert result.eligible_document_count == 1
    assert result.recognized_bom_table_count == 0
    assert result.component_count == 0
    assert result.payload["sources"][0]["tables"] == []


def test_pdf_non_bom_table_remains_not_applicable(monkeypatch) -> None:
    import litoral_trace.us_lacey.product_intelligence_snapshot as module

    parsed = ParsedDocument(
        file_kind="PDF",
        text="Commercial invoice",
        tables=(
            ParsedTable(
                name="page_1_table_1",
                headers=("Invoice", "Amount", "Currency"),
                rows=({"Invoice": "INV-1", "Amount": "100", "Currency": "USD"},),
                source=SourceLocation(
                    page=1,
                    row=1,
                    locator="pdf:page:1;table:1;header_row:1",
                ),
                row_numbers=(2,),
            ),
        ),
        metadata={"page_count": 1},
        ocr_required=False,
    )

    monkeypatch.setattr(module, "parse_document", lambda _filename, _content: parsed)

    result = analyze_product_intelligence_documents(
        (_source("invoice.pdf", b"%PDF-invoice-table"),)
    )

    assert result.status == "NOT_APPLICABLE"
    assert result.eligible_document_count == 1
    assert result.recognized_bom_table_count == 0
    assert result.component_count == 0
