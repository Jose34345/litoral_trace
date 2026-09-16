from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace
from uuid import UUID
import xml.etree.ElementTree as ET

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from litoral_trace.routers import operations as operations_router
from litoral_trace.us_lacey.exporters import (
    LaceyExportHeader,
    LaceyExportPlantLine,
    LaceyExportSnapshot,
    build_lacey_excel,
    build_lawgs_xml,
    consolidate_export_snapshot,
)
from litoral_trace.us_lacey.semantic_evidence_read import EvidenceTextView


OPERATION_ID = "11111111-2222-3333-4444-555555555555"


def _phase_e_snapshot() -> LaceyExportSnapshot:
    return LaceyExportSnapshot(
        operation_id=OPERATION_ID,
        client_reference="PO-1438",
        header=LaceyExportHeader(
            entry_number="123-4567890-1",
            importer="Wood Brokerage International",
            estimated_date_of_arrival="2026-09-20",
        ),
        plant_lines=(
            LaceyExportPlantLine(
                line_reference="LINE-1",
                hts_number="440711",
                entered_value="18600",
                article_component="Sawn pine wood boards",
                genus="Pinus",
                species="taeda",
                country_of_harvest="Argentina",
                quantity="24.5",
                unit="m3",
            ),
        ),
    )


def test_consolidated_export_snapshot_prefers_phase_d_display_text():
    component_field = SimpleNamespace(
        line_reference="LINE-1",
        field_name="article_component",
        scope="PLANT_LINE",
        effective_value="Tablas de madera aserrada de pino",
        proposed_value="Tablas de madera aserrada de pino",
        source_assurance_document_id=7,
        source_page=3,
        source_locator="page:3:block:2",
    )
    hts_field = SimpleNamespace(
        line_reference="LINE-1",
        field_name="hts_code",
        scope="PLANT_LINE",
        effective_value="440711",
        proposed_value="440711",
        source_assurance_document_id=7,
        source_page=3,
        source_locator="page:3:block:3",
    )
    detail = SimpleNamespace(
        public_id=UUID(OPERATION_ID),
        client_reference="PO-1438",
        importer_name="Wood Brokerage International",
        fields=(component_field, hts_field),
        plant_declarations=(),
    )
    translated = EvidenceTextView(
        source_span_id=42,
        target_field="article_component",
        original_text="Tablas de madera aserrada de pino",
        original_language="es",
        translated_text="Sawn pine wood boards",
        display_text="Sawn pine wood boards",
        is_translated=True,
        original_language_label="Spanish",
        source_assurance_document_id=7,
        source_page=3,
        source_locator="page:3:block:2",
    )

    snapshot = consolidate_export_snapshot(
        detail=detail,
        evidence_snapshot={"article_component": (translated,)},
    )

    assert snapshot.plant_lines[0].article_component == "Sawn pine wood boards"
    assert "Tablas de madera" not in snapshot.plant_lines[0].article_component


def test_lawgs_xml_and_excel_share_same_english_projection():
    snapshot = _phase_e_snapshot()

    xml_data = build_lawgs_xml(snapshot)
    root = ET.fromstring(xml_data)
    assert root.tag == "LaceyActDeclaration"
    assert root.findtext("./Header/EntryNumber") == "123-4567890-1"
    assert root.findtext("./PlantLines/PlantLine/ArticleComponent") == "Sawn pine wood boards"
    assert root.findtext("./PlantLines/PlantLine/PlantScientificName/Genus") == "Pinus"

    excel_io = build_lacey_excel(snapshot)
    workbook = load_workbook(excel_io, read_only=True)
    sheet = workbook["Lacey Act Declaration"]
    assert sheet["B1"].value == "123-4567890-1"
    assert sheet["D8"].value == "Sawn pine wood boards"
    assert sheet["E8"].value == "Pinus"


def test_export_endpoints_return_valid_xml_and_xlsx_with_english_text(monkeypatch):
    snapshot = _phase_e_snapshot()
    monkeypatch.setattr(
        operations_router,
        "_export_snapshot",
        lambda *, operation_id, us_session: snapshot,
    )
    app = FastAPI()
    app.include_router(operations_router.router)
    client = TestClient(app)

    xml_response = client.get(f"/operations/{OPERATION_ID}/export/lawgs-xml")
    assert xml_response.status_code == 200
    assert xml_response.headers["content-type"].startswith("application/xml")
    assert f"lawgs_declaration_{OPERATION_ID}.xml" in xml_response.headers["content-disposition"]
    root = ET.fromstring(xml_response.content)
    assert root.findtext("./PlantLines/PlantLine/ArticleComponent") == "Sawn pine wood boards"

    excel_response = client.get(f"/operations/{OPERATION_ID}/export/excel")
    assert excel_response.status_code == 200
    assert excel_response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert f"lacey_summary_{OPERATION_ID}.xlsx" in excel_response.headers["content-disposition"]
    workbook = load_workbook(BytesIO(excel_response.content), read_only=True)
    sheet = workbook["Lacey Act Declaration"]
    assert sheet["D8"].value == "Sawn pine wood boards"
