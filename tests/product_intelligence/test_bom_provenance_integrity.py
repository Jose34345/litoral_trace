from io import BytesIO

from openpyxl import Workbook

from litoral_trace.assurance.parsers import parse_csv, parse_xlsx
from litoral_trace.product_intelligence.bom_ingestion import ingest_bom_table


def test_csv_bom_preserves_physical_row_after_discarded_blank_and_total_rows():
    payload = (
        b"SKU,Product,Component,Material,Qty,Weight,UOM\n"
        b"SKU-A,Chair,Leg,Oak,4,1,kg\n"
        b"\n"
        b"TOTAL,,,,,,\n"
        b"SKU-B,Table,Top,Oak,1,2,kg\n"
    )

    parsed = parse_csv(payload)
    table = parsed.tables[0]

    assert table.row_numbers == (2, 5)

    result = ingest_bom_table(table, document_id="doc-csv")
    by_sku = {item.sku: item for item in result.compositions}
    assert by_sku["SKU-A"].components[0].source.row == 2
    assert by_sku["SKU-B"].components[0].source.row == 5
    assert by_sku["SKU-B"].components[0].component_key == "SKU-B:row:5"


def test_xlsx_bom_preserves_physical_row_after_discarded_blank_row():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "BOM"
    sheet.append(["SKU", "Product", "Component", "Material", "Qty", "Weight", "UOM"])
    sheet.append(["SKU-A", "Chair", "Leg", "Oak", "4", "1", "kg"])
    sheet.append([None, None, None, None, None, None, None])
    sheet.append(["SKU-B", "Table", "Top", "Oak", "1", "2", "kg"])
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()

    parsed = parse_xlsx(buffer.getvalue())
    table = parsed.tables[0]

    assert table.row_numbers == (2, 4)
    result = ingest_bom_table(table)
    by_sku = {item.sku: item for item in result.compositions}
    assert by_sku["SKU-B"].components[0].source.row == 4


def test_conflicting_product_name_for_same_sku_is_reported_and_row_is_not_merged():
    payload = (
        b"SKU,Product,Component,Material,Qty,Weight,UOM\n"
        b"SKU-A,Chair,Leg,Oak,4,1,kg\n"
        b"SKU-A,Table,Top,Oak,1,2,kg\n"
    )

    table = parse_csv(payload).tables[0]
    result = ingest_bom_table(table)

    assert len(result.compositions) == 1
    composition = result.compositions[0]
    assert composition.sku == "SKU-A"
    assert composition.product_name == "Chair"
    assert [component.description_raw for component in composition.components] == ["Leg"]
    assert [issue.code for issue in result.issues] == ["CONFLICTING_PRODUCT_NAME"]
    assert result.issues[0].source.row == 3
