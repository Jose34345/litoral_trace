from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook

from litoral_trace.assurance.parsers import parse_csv, parse_xlsx
from litoral_trace.product_intelligence.bom_ingestion import ingest_bom_table


def test_csv_bom_integration_preserves_sku_isolation_and_units():
    content = (
        "SKU,Product,Component,Material,Qty,Weight,UOM\n"
        "SKU-A,Chair,Leg,Oak,4,500,g\n"
        "SKU-A,Chair,Seat,MDF,1,1,kg\n"
        "SKU-B,Tool,Handle,Beech,1,2,lb\n"
    ).encode("utf-8")

    parsed = parse_csv(content)
    result = ingest_bom_table(parsed.tables[0], document_id="csv-doc")

    assert result.issues == ()
    by_sku = {item.sku: item for item in result.compositions}
    assert set(by_sku) == {"SKU-A", "SKU-B"}
    assert by_sku["SKU-A"].components[0].material.mass.kilograms == Decimal("0.500")
    assert by_sku["SKU-B"].components[0].material.mass.kilograms == Decimal("0.90718474")
    assert by_sku["SKU-A"].components[0].source.row == 2


def test_xlsx_bom_integration_uses_existing_parser_and_sheet_provenance():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Components"
    sheet.append(["Item Number", "Product", "Part", "Material Description", "Qty", "Weight", "UOM"])
    sheet.append(["SKU-X", "Desk", "Top", "Oak", 1, 4, "lb"])
    sheet.append(["SKU-X", "Desk", "Frame", "Steel", 1, 2, "kg"])
    sheet.append(["SKU-Y", "Tray", "Base", "Plywood", 1, 4, "oz"])
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()

    parsed = parse_xlsx(stream.getvalue())
    result = ingest_bom_table(parsed.tables[0], document_id="xlsx-doc")

    assert result.issues == ()
    by_sku = {item.sku: item for item in result.compositions}
    assert set(by_sku) == {"SKU-X", "SKU-Y"}
    first = by_sku["SKU-X"].components[0]
    assert first.source.document_id == "xlsx-doc"
    assert first.source.sheet == "Components"
    assert first.source.row == 2
    assert first.material.mass.kilograms == Decimal("1.81436948")


def test_valid_rows_survive_when_an_independent_bom_row_is_incomplete():
    content = (
        "SKU,Product,Component,Material,Qty,Weight,UOM\n"
        "GOOD,Chair,Leg,Oak,1,1,kg\n"
        "BAD,Chair,Seat,,1,1,kg\n"
    ).encode("utf-8")

    parsed = parse_csv(content)
    result = ingest_bom_table(parsed.tables[0])

    assert [item.sku for item in result.compositions] == ["GOOD"]
    assert [issue.code for issue in result.issues] == ["MISSING_MATERIAL"]
